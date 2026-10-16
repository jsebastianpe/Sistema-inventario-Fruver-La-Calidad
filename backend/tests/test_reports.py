"""Los reportes deben coincidir con la información almacenada."""
import io
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy import text

from app.core.timeutil import today_local
from tests.helpers import category, key, move, product


def scenario(owner, db):
    cat = category(owner)
    sup = owner.post("/api/v1/proveedores", json={"name": "=Proveedor peligroso"}).json()["id"]
    a = product(owner, db, sku="A", name="Aguacate", cat=cat, price="5000", min_stock="50")
    b = product(owner, db, sku="B", name="Banano", cat=cat, price="2000")
    pid = owner.post("/api/v1/compras", json={
        "supplier_id": sup, "purchase_date": today_local().isoformat(),
        "lines": [{"product_id": a["id"], "quantity": "40", "unit_price": "3000", "waste_qty": "2",
                   "expiry_date": today_local().isoformat()},
                  {"product_id": b["id"], "quantity": "100", "unit_price": "1000"}]}).json()["id"]
    owner.post(f"/api/v1/compras/{pid}/confirmar")
    for lines in ([{"product_id": a["id"], "quantity": "3"}],
                  [{"product_id": b["id"], "quantity": "10"}, {"product_id": a["id"], "quantity": "1"}]):
        r = owner.post("/api/v1/ventas", json={"lines": lines}, headers={"Idempotency-Key": key()})
        assert r.status_code == 201
    s = owner.post("/api/v1/ventas", json={"lines": [{"product_id": b["id"], "quantity": "5"}]},
                   headers={"Idempotency-Key": key()}).json()
    owner.post(f"/api/v1/ventas/{s['id']}/anular", json={"reason": "anulada para prueba"})
    move(owner, b["id"], "merma", "4", motivo="Maduración")
    return a, b


def test_sales_report_matches_database(owner, db):
    scenario(owner, db)
    r = owner.get("/api/v1/reportes/ventas-por-periodo").json()
    db_total = db.execute(text("SELECT sum(total), sum(cost_total), count(*) FROM sales WHERE status='confirmada'")).one()
    assert Decimal(r["totales"]["total"]) == db_total[0] == Decimal("40000.00")   # 4·5000 + 10·2000
    assert Decimal(r["totales"]["costo_vendido"]) == db_total[1]
    assert r["totales"]["ventas"] == db_total[2] == 2                              # excluye la anulada
    assert any("NO es utilidad neta" in n for n in r["notas"])


def test_margin_report_consistent_with_sales(owner, db):
    scenario(owner, db)
    r = owner.get("/api/v1/reportes/margen-bruto").json()
    v = owner.get("/api/v1/reportes/ventas-por-periodo").json()
    assert r["totales"]["margen_bruto"] == v["totales"]["margen_bruto"]
    assert r["titulo"].startswith("Margen bruto estimado")


def test_stock_report_and_value(owner, db):
    a, b = scenario(owner, db)
    r = owner.get("/api/v1/reportes/existencias").json()
    rows = {x["sku"]: x for x in r["filas"]}
    assert rows["A"]["existencia"] == "34.0000"          # 40 − 2 merma recepción − 4 vendidos
    assert rows["A"]["bajo_minimo"] == "Sí"
    assert rows["B"]["existencia"] == "86.0000"          # 100 − 10 − 4 (la venta anulada se devolvió)
    expected = db.execute(text("SELECT round(sum(greatest(quantity,0) * avg_cost), 2) FROM stock_balances")).scalar()
    assert Decimal(r["totales"]["valor_estimado"]) == expected


def test_purchases_and_waste_reports(owner, db):
    scenario(owner, db)
    p = owner.get("/api/v1/reportes/compras-por-proveedor").json()
    assert p["filas"][0]["total"] == "220000.00" and p["filas"][0]["compras"] == 1
    w = {x["producto"]: x for x in owner.get("/api/v1/reportes/mermas").json()["filas"]}
    assert w["Banano"]["merma_inventario"] == "4.0000" and w["Aguacate"]["merma_recepcion"] == "2.0000"


def test_dashboard_and_alerts(owner, employee, db):
    scenario(owner, db)
    d = owner.get("/api/v1/panel").json()
    assert d["productos_activos"] == 2
    assert d["bajo_minimo"]["total"] == 1 and d["por_vencer"]["total"] == 1
    assert d["ventas"]["cantidad"] == 2 and d["ventas"]["total"] == "40000.00"
    assert d["compras"]["total"] == "220000.00"
    assert d["mayor_movimiento"][0]["producto"] == "Banano"
    assert "valor_inventario_estimado" in d
    e = employee.get("/api/v1/panel").json()
    assert "ventas" not in e and "valor_inventario_estimado" not in e and "compras" not in e
    alerts = owner.get("/api/v1/alertas").json()
    kinds = {a["tipo"] for a in alerts}
    assert {"existencia_baja", "vencimiento"} <= kinds


def test_report_permissions(owner, admin, employee, db):
    scenario(owner, db)
    assert employee.get("/api/v1/reportes/existencias").status_code == 200
    assert "costo_promedio" not in employee.get("/api/v1/reportes/existencias").json()["filas"][0]
    for name in ("ventas-por-periodo", "margen-bruto", "compras-por-proveedor", "auditoria", "costo-vs-mercado"):
        assert employee.get(f"/api/v1/reportes/{name}").status_code == 403, name
    assert admin.get("/api/v1/reportes/auditoria").status_code == 403
    assert owner.get("/api/v1/reportes/auditoria").status_code == 200
    assert employee.get("/api/v1/reportes/existencias", params={"formato": "csv"}).status_code == 403
    names = {r["nombre"] for r in employee.get("/api/v1/reportes").json()}
    assert names == {"existencias", "movimientos", "rotacion", "mermas", "evolucion-mercado"}


def test_csv_and_xlsx_exports(owner, db):
    scenario(owner, db)
    r = owner.get("/api/v1/reportes/compras-por-proveedor", params={"formato": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    body = r.content.decode("utf-8-sig")
    assert body.splitlines()[0].startswith("Proveedor,") and "220000.00" in body
    r = owner.get("/api/v1/reportes/compras-por-proveedor", params={"formato": "xlsx"})
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert "Compras por proveedor" in values and 220000.0 in values
    assert "'=Proveedor peligroso" in values      # sin inyección de fórmulas
    logged = db.execute(text("SELECT count(*) FROM audit_log WHERE action='reportes.exportar'")).scalar()
    assert logged == 2


def test_invalid_range(owner):
    r = owner.get("/api/v1/reportes/movimientos", params={"desde": "2026-10-10", "hasta": "2026-10-01"})
    assert r.status_code == 400
    assert owner.get("/api/v1/reportes/inventado").status_code == 404
    assert owner.get("/api/v1/reportes/evolucion-costos").status_code == 400   # requiere producto
