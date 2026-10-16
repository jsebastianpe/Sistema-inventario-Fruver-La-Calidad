"""Exportación de reportes a CSV y Excel (XLSX)."""
import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.timeutil import local_tz, now_utc
from app.domains.reports.service import Report

NUMBER_FORMATS = {
    "dinero": '"$" #,##0.00',
    "cantidad": "#,##0.0000",
    "numero": "#,##0.##",
    "porcentaje": "0.00%",
}


def _plain(v):
    if v is None:
        return ""
    if isinstance(v, Decimal):
        return format(v, "f")
    if isinstance(v, date | datetime):
        return v.isoformat()
    return v


def to_csv(report: Report) -> bytes:
    """CSV estándar (coma como separador, punto decimal) con BOM UTF-8 para Excel.

    Para abrir con formato numérico colombiano se recomienda la exportación XLSX.
    """
    buf = io.StringIO()
    w = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
    w.writerow([c.title for c in report.columns])
    for row in report.rows:
        w.writerow([_plain(row.get(c.key)) for c in report.columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _safe_cell(v):
    """Evita inyección de fórmulas en hojas de cálculo (=, +, -, @ al inicio de un texto)."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def to_xlsx(report: Report, user_name: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    generated = now_utc().astimezone(local_tz()).strftime("%Y-%m-%d %H:%M")
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generado: {generated} (America/Bogota) por {user_name}"])
    if report.filters:
        ws.append(["Filtros: " + ", ".join(f"{k}={_plain(v)}" for k, v in report.filters.items())])
    for note in report.notes:
        ws.append([f"Nota: {note}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.title for c in report.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6B3A")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in report.rows:
        values = []
        for c in report.columns:
            v = row.get(c.key)
            if isinstance(v, Decimal):
                v = float(v)  # Excel solo maneja números binarios; el CSV conserva el valor decimal exacto
            values.append(_safe_cell(v))
        ws.append(values)
        r = ws.max_row
        for i, c in enumerate(report.columns, 1):
            if c.kind in NUMBER_FORMATS:
                ws.cell(row=r, column=i).number_format = NUMBER_FORMATS[c.kind]
    if report.totals:
        ws.append([])
        ws.append(["Totales"])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for k, v in report.totals.items():
            col = next((c for c in report.columns if c.key == k), None)
            ws.append([col.title if col else k, float(v) if isinstance(v, Decimal) else _plain(v)])
            if col and col.kind in NUMBER_FORMATS:
                ws.cell(row=ws.max_row, column=2).number_format = NUMBER_FORMATS[col.kind]
    for i, c in enumerate(report.columns, 1):
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = max(12, min(40, len(c.title) + 4))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
